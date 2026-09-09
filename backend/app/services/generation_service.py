"""Server-Side Requirement Document Generator (DOCX & PDF).

Generates official 9-section MSME Client Requirement Documents as durable
Word (.docx) and PDF artifacts.
"""
from __future__ import annotations

import os
import uuid
import datetime
from typing import Any, Dict, Optional, Tuple
from sqlalchemy.ext.asyncio import AsyncSession

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from app.core.config import settings
from app.models.domain import Application, GeneratedDocument
from app.services.template_service import TemplateService


class GenerationService:
    """Renders requirement documents to DOCX and PDF and persists artifact metadata."""

    @staticmethod
    def _ensure_dir(application_id: str) -> str:
        out_dir = os.path.join(settings.STORAGE_DIR, "generated", str(application_id))
        os.makedirs(out_dir, exist_ok=True)
        return out_dir

    @staticmethod
    def generate_docx(template_data: Dict[str, Any], output_path: str) -> str:
        """Render standard 9-section template into a styled Microsoft Word (.docx) file."""
        doc = docx.Document()

        # Page margins
        sections = doc.sections
        for s in sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Header Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(template_data.get("header", {}).get("system_title", "UDYAM REGISTRATION"))
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(30, 27, 75)

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(template_data.get("header", {}).get("form_subtitle", "Client Information & Requirement Form"))
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = RGBColor(79, 70, 229)

        # Meta Banner
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        b_type = template_data.get("header", {}).get("business_type", "General Enterprise")
        o_type = template_data.get("header", {}).get("org_type", "Proprietorship")
        meta_run = meta_p.add_run(f"Business Type: {b_type}   |   Organisation Type: {o_type}")
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # Sections
        for sec in template_data.get("sections", []):
            # Section Header Heading
            h_p = doc.add_paragraph()
            h_run = h_p.add_run(sec.get("section_title", "SECTION"))
            h_run.bold = True
            h_run.font.size = Pt(11)
            h_run.font.color.rgb = RGBColor(49, 46, 129)
            h_p.paragraph_format.space_before = Pt(12)
            h_p.paragraph_format.space_after = Pt(4)

            # Table for Fields
            fields = sec.get("fields", [])
            if fields:
                table = doc.add_table(rows=1, cols=2)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                # Set header row
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = "Information Required"
                hdr_cells[1].text = "Client Details"
                hdr_cells[0].paragraphs[0].runs[0].bold = True
                hdr_cells[1].paragraphs[0].runs[0].bold = True
                hdr_cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
                hdr_cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)

                for f in fields:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(f.get("label", ""))
                    val = str(f.get("value") or "")
                    row_cells[1].text = val if val.strip() else "________________________________________"
                    row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)
                    row_cells[1].paragraphs[0].runs[0].font.size = Pt(9)

            # Documents Required
            docs_req = sec.get("documents_required", [])
            if docs_req:
                dp = doc.add_paragraph()
                dp.add_run("Documents Required:").bold = True
                dp.paragraph_format.space_before = Pt(4)
                for d in docs_req:
                    dp2 = doc.add_paragraph(f"  [ {'X' if d.get('checked') else ' '} ]  {d.get('label', '')}")
                    dp2.paragraph_format.space_after = Pt(2)

            # Checkboxes / Declarations
            checkboxes = sec.get("checkboxes", [])
            if checkboxes:
                for chk in checkboxes:
                    cp = doc.add_paragraph()
                    cp.paragraph_format.space_before = Pt(4)
                    c_run = cp.add_run(f"[ X ] {chk.get('label', '')}: ")
                    c_run.bold = True
                    cp.add_run(str(chk.get("text", "")))

        doc.save(output_path)
        return output_path

    @staticmethod
    def generate_pdf(template_data: Dict[str, Any], output_path: str) -> str:
        """Render standard 9-section template into a clean vector PDF using ReportLab."""
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E1B4B"),
            alignment=1,
            spaceAfter=2,
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#4F46E5"),
            alignment=1,
            spaceAfter=6,
        )
        meta_style = ParagraphStyle(
            "DocMeta",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748B"),
            alignment=1,
            spaceAfter=14,
        )
        sec_title_style = ParagraphStyle(
            "SecTitle",
            parent=styles["Heading2"],
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#312E81"),
            spaceBefore=10,
            spaceAfter=4,
        )
        cell_style = ParagraphStyle(
            "CellText",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#1E293B"),
        )
        cell_bold = ParagraphStyle(
            "CellBold",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#0F172A"),
            fontName="Helvetica-Bold",
        )

        story = []

        # Header
        sys_title = template_data.get("header", {}).get("system_title", "UDYAM REGISTRATION")
        form_sub = template_data.get("header", {}).get("form_subtitle", "Client Information & Requirement Form")
        b_type = template_data.get("header", {}).get("business_type", "General Enterprise")
        o_type = template_data.get("header", {}).get("org_type", "Proprietorship")

        story.append(Paragraph(sys_title, title_style))
        story.append(Paragraph(form_sub, subtitle_style))
        story.append(Paragraph(f"<b>Business Type:</b> {b_type} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Organisation Type:</b> {o_type}", meta_style))

        for sec in template_data.get("sections", []):
            story.append(Paragraph(sec.get("section_title", ""), sec_title_style))

            fields = sec.get("fields", [])
            if fields:
                table_data = [
                    [Paragraph("<b>Information Required</b>", cell_bold), Paragraph("<b>Client Details</b>", cell_bold)]
                ]
                for f in fields:
                    val = str(f.get("value") or "")
                    disp_val = val if val.strip() else "________________________________________"
                    table_data.append([
                        Paragraph(str(f.get("label", "")), cell_style),
                        Paragraph(disp_val, cell_style),
                    ])

                t = Table(table_data, colWidths=[200, 320])
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.append(t)
                story.append(Spacer(1, 4))

            docs_req = sec.get("documents_required", [])
            if docs_req:
                story.append(Paragraph("<b>Documents Required:</b>", cell_bold))
                for d in docs_req:
                    mark = "☒" if d.get("checked") else "☐"
                    story.append(Paragraph(f"&nbsp;&nbsp;{mark} {d.get('label', '')}", cell_style))
                story.append(Spacer(1, 4))

            checkboxes = sec.get("checkboxes", [])
            if checkboxes:
                for chk in checkboxes:
                    story.append(Paragraph(f"<b>[ X ] {chk.get('label', '')}:</b> {chk.get('text', '')}", cell_style))
                story.append(Spacer(1, 4))

        doc.build(story)
        return output_path

    @staticmethod
    async def generate_requirement_documents(
        application_id: uuid.UUID,
        db: AsyncSession,
        user_id: Optional[uuid.UUID] = None,
        custom_template_data: Optional[Dict[str, Any]] = None,
    ) -> Tuple[GeneratedDocument, GeneratedDocument]:
        """Generate both PDF and DOCX requirement artifacts and save to database."""
        out_dir = GenerationService._ensure_dir(str(application_id))

        if custom_template_data:
            template_data = custom_template_data
        else:
            template_data = await TemplateService.build_template_for_application(application_id, db)

        file_stem = f"Requirement_Doc_{str(application_id)[:8]}_{int(datetime.datetime.now().timestamp())}"
        pdf_path = os.path.join(out_dir, f"{file_stem}.pdf")
        docx_path = os.path.join(out_dir, f"{file_stem}.docx")

        # Generate files
        GenerationService.generate_pdf(template_data, pdf_path)
        GenerationService.generate_docx(template_data, docx_path)

        pdf_size = os.path.getsize(pdf_path) if os.path.exists(pdf_path) else 0
        docx_size = os.path.getsize(docx_path) if os.path.exists(docx_path) else 0

        title = template_data.get("template_name") or "MSME Requirement Document"

        # Record GeneratedDocument entries
        pdf_record = GeneratedDocument(
            application_id=application_id,
            document_title=f"{title} (PDF)",
            format="PDF",
            storage_path=pdf_path,
            file_size_bytes=pdf_size,
            status="GENERATED",
            generated_by=user_id,
        )
        docx_record = GeneratedDocument(
            application_id=application_id,
            document_title=f"{title} (Word)",
            format="DOCX",
            storage_path=docx_path,
            file_size_bytes=docx_size,
            status="GENERATED",
            generated_by=user_id,
        )

        db.add(pdf_record)
        db.add(docx_record)
        await db.commit()
        await db.refresh(pdf_record)
        await db.refresh(docx_record)

        return pdf_record, docx_record
